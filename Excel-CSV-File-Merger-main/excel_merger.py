"""
Excel & CSV File Merger
=======================
Scans a folder for all Excel (.xlsx, .xls) and CSV files, intelligently
merges them into a single normalized master file, and outputs a clean Excel
file. Handles messy real-world data: different column orders, missing columns,
duplicate rows, inconsistent formatting, and varied column naming conventions.

Usage:
    python excel_merger.py --input ./sample_input --output merged_master.xlsx
    python excel_merger.py --input /path/to/folder --no-dedup

Features:
    - Discovers all .xlsx, .xls, and .csv files recursively
    - Normalizes column names (strips whitespace, lowercases for matching)
    - Maps common column aliases (e.g. "Sales Amount" → "revenue")
    - Fills missing columns with empty values
    - Deduplicates identical rows (configurable)
    - Strips leading/trailing whitespace from all text cells
    - Normalizes currency strings ("$5,400.00" → 5400.00)
    - Adds a "source_file" column tracking which file each row came from
    - Outputs a formatted Excel file with statistics summary sheet
"""

import argparse
import logging
import re
from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill

logger = logging.getLogger(__name__)

SUPPORTED_SUFFIXES = {".xlsx", ".xls", ".csv"}
SOURCE_COLUMN = "source_file"


# ---------------------------------------------------------------------------
# Logging
# ---------------------------------------------------------------------------
def configure_logging(log_path: Path = Path("merger.log")) -> None:
    """Configure console and file logging for the command-line entry point.

    Logging is deliberately configured at runtime rather than import time so
    importing this module as a library neither creates a file nor fails merely
    because the current directory is read-only.
    """
    handlers: list[logging.Handler] = [logging.StreamHandler()]
    file_error: OSError | None = None
    try:
        # Input files are arbitrary user data, so the log is explicitly UTF-8
        # rather than the platform default (cp1252 on Windows).
        handlers.append(logging.FileHandler(log_path, encoding="utf-8"))
    except OSError as exc:
        file_error = exc

    logging.basicConfig(
        level=logging.INFO,
        format="%(asctime)s [%(levelname)s] %(message)s",
        handlers=handlers,
        force=True,
    )
    if file_error is not None:
        logger.warning("Could not write log file %s: %s", log_path, file_error)

# ---------------------------------------------------------------------------
# Column alias map — maps known variants to a canonical name
# Add more entries here as you encounter new file formats.
# ---------------------------------------------------------------------------
COLUMN_ALIASES: dict[str, str] = {
    # Name variants
    "full name": "name",
    "full_name": "name",
    "first name": "name",
    "employee name": "name",
    # Email variants
    "contact email": "email",
    "email address": "email",
    "e-mail": "email",
    # Revenue / money variants
    "revenue": "revenue",
    "sales amount": "revenue",
    "sales_amount": "revenue",
    "amount": "revenue",
    "salary": "revenue",
    "total": "revenue",
    # Region / territory variants
    "region": "region",
    "territory": "region",
    "area": "region",
    "zone": "region",
    # Date variants
    "date": "date",
    "transaction date": "date",
    "transaction_date": "date",
    "hire date": "date",
    "hire_date": "date",
    "sale date": "date",
    # Department variants
    "department": "department",
    "dept": "department",
    "team": "department",
    "division": "department",
}


def normalize_column_name(col: str) -> str:
    """
    Lowercase, strip whitespace, and apply alias mapping.
    Returns the canonical column name.
    """
    cleaned = str(col).strip().lower()
    return COLUMN_ALIASES.get(cleaned, cleaned)


def dedupe_columns(cols: list[str], source_filename: str) -> list[str]:
    """
    Ensure canonical column names are unique within a single file.

    Multiple source headers can alias to the same canonical name (e.g. both
    "Amount" and "Total" map to "revenue"). Left unhandled, pandas produces two
    identically-named columns, `df["revenue"]` returns a DataFrame instead of a
    Series, and every downstream scalar operation raises
    "The truth value of a Series is ambiguous".
    """
    used: set[str] = set()
    next_suffix: dict[str, int] = {}
    out: list[str] = []
    for c in cols:
        renamed = c
        if renamed in used:
            suffix = next_suffix.get(c, 2)
            renamed = f"{c}_{suffix}"
            while renamed in used:
                suffix += 1
                renamed = f"{c}_{suffix}"
            next_suffix[c] = suffix + 1
            logger.warning(
                f"{source_filename}: duplicate canonical column '{c}' renamed to "
                f"'{renamed}'. Review COLUMN_ALIASES if these should be merged."
            )
        used.add(renamed)
        next_suffix.setdefault(c, 2)
        out.append(renamed)
    return out


def normalize_currency(value) -> float | None:
    """
    Convert currency strings like '$5,400.00' or '5400' to float.
    Returns None if conversion is not possible.
    """
    if isinstance(value, (pd.Series, pd.DataFrame)):
        raise TypeError(
            "normalize_currency expects a scalar. Receiving a Series means the "
            "DataFrame has duplicate column names -- run dedupe_columns() first."
        )
    if pd.isna(value):
        return None
    raw = str(value).strip()
    # Remove currency symbols and thousands separators
    cleaned = re.sub(r"[^\d.\-]", "", raw)
    try:
        return float(cleaned) if cleaned else None
    except ValueError:
        return None


def clean_string(value) -> str:
    """Strip whitespace from string values; leave non-strings unchanged."""
    if isinstance(value, str):
        return value.strip()
    return value


# ---------------------------------------------------------------------------
# File loading
# ---------------------------------------------------------------------------
def load_file(path: Path) -> pd.DataFrame | None:
    """
    Load an Excel or CSV file into a DataFrame.
    Returns None if the file cannot be read.
    """
    try:
        suffix = path.suffix.lower()
        if suffix in (".xlsx", ".xls"):
            df = pd.read_excel(path, engine="openpyxl" if suffix == ".xlsx" else "xlrd")
        elif suffix == ".csv":
            df = pd.read_csv(path)
        else:
            logger.warning(f"Unsupported file type: {path.name}")
            return None

        logger.info(f"Loaded {path.name}: {len(df)} rows, {len(df.columns)} columns")
        return df

    except Exception as e:
        logger.error(f"Failed to load {path.name}: {e}")
        return None


# ---------------------------------------------------------------------------
# Per-file normalization
# ---------------------------------------------------------------------------
def normalize_dataframe(df: pd.DataFrame, source_filename: str) -> pd.DataFrame:
    """
    Apply all normalization steps to a single DataFrame:
    1. Normalize column names (strip + alias mapping)
    2. Strip whitespace from string cells
    3. Drop entirely empty rows
    4. Normalize currency values in 'revenue' column
    5. Add source_file column
    """
    # Work on a copy so callers do not get partially mutated input if a later
    # normalization step fails.
    df = df.copy()

    # --- 1. Normalize column names ---
    # Reserve SOURCE_COLUMN before normalizing user columns. Otherwise an input
    # column named "source_file" would be silently overwritten in step 5.
    normalized_columns = [normalize_column_name(c) for c in df.columns]
    df.columns = dedupe_columns(
        [SOURCE_COLUMN, *normalized_columns], source_filename
    )[1:]

    # --- 2. Strip whitespace from all string cells ---
    df = df.map(clean_string)

    # --- 3. Drop rows where ALL values are missing or empty after trimming ---
    # dropna alone does not consider "" empty, so a row made entirely of
    # whitespace survived step 2 in earlier versions.
    empty_cells = df.isna() | df.eq("")
    df = df.loc[~empty_cells.all(axis=1)].copy()

    # --- 4. Normalize revenue column ---
    if "revenue" in df.columns:
        df["revenue"] = df["revenue"].apply(normalize_currency)

    # --- 5. Tag with source ---
    df[SOURCE_COLUMN] = source_filename

    return df


# ---------------------------------------------------------------------------
# Merge engine
# ---------------------------------------------------------------------------
def merge_files(
    input_folder: Path,
    output_path: Path,
    deduplicate: bool = True,
) -> None:
    """
    Main merge function. Discovers files, normalizes each, aligns columns,
    concatenates, deduplicates, and writes to Excel.
    """
    # --- Discover all supported files ---
    # Inspect suffixes case-insensitively so names such as DATA.CSV work on
    # case-sensitive platforms too. The requested output is excluded so a
    # rerun cannot ingest its own previous result.
    output_resolved = output_path.resolve()
    found_files = sorted(
        (
            path
            for path in input_folder.rglob("*")
            if path.is_file()
            and path.suffix.lower() in SUPPORTED_SUFFIXES
            and not path.name.startswith("~$")
            and path.resolve() != output_resolved
        ),
        key=lambda path: path.relative_to(input_folder).as_posix().casefold(),
    )

    if not found_files:
        raise FileNotFoundError(f"No Excel or CSV files found in: {input_folder}")

    discovered_names = [path.relative_to(input_folder).as_posix() for path in found_files]
    logger.info(f"Found {len(found_files)} file(s) to merge: {discovered_names}")

    # --- Load and normalize each file ---
    frames: list[pd.DataFrame] = []
    load_stats: list[dict] = []

    for path in found_files:
        source_filename = path.relative_to(input_folder).as_posix()
        df = load_file(path)
        if df is None:
            continue
        original_rows = len(df)
        df = normalize_dataframe(df, source_filename)
        frames.append(df)
        load_stats.append({
            "file": source_filename,
            "rows_loaded": original_rows,
            "columns": len(df.columns) - 1,  # exclude source_file
        })

    if not frames:
        raise ValueError("No files could be successfully loaded")

    # --- Concatenate (pandas aligns on column names, fills missing with NaN) ---
    merged = pd.concat(frames, ignore_index=True, sort=False)
    rows_before_dedup = len(merged)
    logger.info(f"Combined total: {rows_before_dedup} rows")

    # --- Deduplicate (exclude source_file from key) ---
    dupes_removed = 0
    if deduplicate:
        key_cols = [c for c in merged.columns if c != SOURCE_COLUMN]
        merged.drop_duplicates(subset=key_cols, keep="first", inplace=True)
        merged.reset_index(drop=True, inplace=True)
        dupes_removed = rows_before_dedup - len(merged)
        if dupes_removed:
            logger.info(f"Removed {dupes_removed} duplicate row(s)")

    # --- Reorder: put source_file last ---
    cols = [c for c in merged.columns if c != SOURCE_COLUMN] + [SOURCE_COLUMN]
    merged = merged[cols]

    # --- Write output ---
    write_excel(merged, output_path, load_stats, dupes_removed)
    logger.info(f"Merge complete: {len(merged)} rows written to {output_path}")
    print(f"\n[OK] Merged {len(frames)} files -> {len(merged)} rows -> {output_path}")


# ---------------------------------------------------------------------------
# Excel writer with formatting
# ---------------------------------------------------------------------------
def write_excel(df: pd.DataFrame, output_path: Path, stats: list[dict], dupes_removed: int) -> None:
    """Write the merged DataFrame plus a summary sheet to a formatted Excel file."""
    output_path.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        # --- Main data sheet ---
        df.to_excel(writer, index=False, sheet_name="Merged Data")
        _format_sheet(writer.sheets["Merged Data"], df)

        # --- Summary sheet ---
        summary_rows = []
        for s in stats:
            summary_rows.append(s)
        summary_rows.append({"file": "─" * 20, "rows_loaded": "─" * 10, "columns": "─" * 10})
        summary_rows.append({
            "file": "TOTAL",
            "rows_loaded": sum(s["rows_loaded"] for s in stats),
            "columns": "N/A",
        })
        summary_rows.append({"file": "Duplicates removed", "rows_loaded": dupes_removed, "columns": ""})
        summary_rows.append({"file": "Final row count", "rows_loaded": len(df), "columns": ""})

        summary_df = pd.DataFrame(summary_rows, columns=["file", "rows_loaded", "columns"])
        summary_df.columns = ["Source File", "Rows Loaded", "Original Columns"]
        summary_df.to_excel(writer, index=False, sheet_name="Merge Summary")
        _format_sheet(writer.sheets["Merge Summary"], summary_df)


def _format_sheet(ws, df: pd.DataFrame) -> None:
    """Apply header styling and auto-fit columns to a worksheet."""
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1A3A5C", end_color="1A3A5C", fill_type="solid")

    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", wrap_text=True)

    # Auto-fit column widths based on content
    for col in ws.columns:
        max_len = max(
            (len(str(cell.value or "")) for cell in col),
            default=10,
        )
        ws.column_dimensions[col[0].column_letter].width = min(max_len + 2, 50)

    # Alternate row shading for readability
    light_fill = PatternFill(start_color="F0F4F8", end_color="F0F4F8", fill_type="solid")
    for row_idx, row in enumerate(ws.iter_rows(min_row=2), start=2):
        if row_idx % 2 == 0:
            for cell in row:
                cell.fill = light_fill


# ---------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------
def main():
    configure_logging()
    parser = argparse.ArgumentParser(description="Excel & CSV File Merger")
    parser.add_argument(
        "--input",
        default="./sample_input",
        help="Folder containing Excel/CSV files to merge (default: ./sample_input)",
    )
    parser.add_argument(
        "--output",
        default="merged_master.xlsx",
        help="Output Excel filename (default: merged_master.xlsx)",
    )
    parser.add_argument(
        "--no-dedup",
        action="store_true",
        help="Skip duplicate row removal",
    )
    args = parser.parse_args()

    input_folder = Path(args.input)
    if not input_folder.is_dir():
        raise ValueError(f"Input path is not a directory: {input_folder}")

    merge_files(
        input_folder=input_folder,
        output_path=Path(args.output),
        deduplicate=not args.no_dedup,
    )


if __name__ == "__main__":
    main()
