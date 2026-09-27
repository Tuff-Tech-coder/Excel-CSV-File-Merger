"""Regression tests for the Excel/CSV merge pipeline."""

import json
import os
import subprocess
import sys
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import Workbook, load_workbook

PROJECT_ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(PROJECT_ROOT))

import excel_merger as merger  # noqa: E402
from create_samples import create_samples  # noqa: E402
from excel_merger import (  # noqa: E402
    COLUMN_ALIASES,
    clean_string,
    dedupe_columns,
    load_file,
    merge_files,
    normalize_column_name,
    normalize_currency,
    normalize_dataframe,
    validate_json_output_path,
    write_json_summary,
)


class TestNormalizeCurrency:
    @pytest.mark.parametrize("raw,expected", [
        ("$5,400.00", 5400.0),
        ("5400", 5400.0),
        ("  USD 1,299.50 ", 1299.5),
        ("-250.75", -250.75),
        ("(1,234.56)", -1234.56),
        (3100, 3100.0),
        (3100.25, 3100.25),
    ])
    def test_parses_supported_currency_forms(self, raw, expected):
        assert normalize_currency(raw) == expected

    @pytest.mark.parametrize("raw", [
        "", "N/A", None, float("nan"), True,
        "1e3", "invoice 123", "1.234,56 EUR",
    ])
    def test_rejects_missing_or_ambiguous_values(self, raw):
        assert normalize_currency(raw) is None

    def test_rejects_series_with_clear_message(self):
        with pytest.raises(TypeError, match="duplicate column names"):
            normalize_currency(pd.Series([1, 2]))


class TestNormalizeColumnName:
    @pytest.mark.parametrize("raw,expected", [
        ("Sales Amount", "revenue"),
        ("  TERRITORY  ", "region"),
        ("Transaction Date", "date"),
        ("Hire Date", "hire_date"),
        ("Salary", "salary"),
        ("Dept", "department"),
        ("Unmapped Column", "unmapped column"),
        ("Sales    Amount", "revenue"),
    ])
    def test_aliases_and_normalization(self, raw, expected):
        assert normalize_column_name(raw) == expected

    def test_alias_map_is_idempotent(self):
        for canonical in set(COLUMN_ALIASES.values()):
            assert normalize_column_name(canonical) == canonical

    def test_formula_header_is_neutralized(self):
        assert normalize_column_name("=2+3") == "'=2+3"


class TestDedupeColumns:
    def test_disambiguates_aliased_collisions(self):
        assert dedupe_columns(["name", "revenue", "revenue"], "f.csv") == [
            "name", "revenue", "revenue_2",
        ]

    def test_leaves_unique_columns_untouched(self):
        columns = ["name", "email", "revenue"]
        assert dedupe_columns(columns, "f.csv") == columns

    def test_handles_triple_collision(self):
        assert dedupe_columns(["r", "r", "r"], "f.csv") == ["r", "r_2", "r_3"]

    def test_generated_name_cannot_collide_with_existing_name(self):
        columns = dedupe_columns(["x", "x", "x_2"], "f.csv")
        assert columns == ["x", "x_2", "x_2_2"]
        assert len(columns) == len(set(columns))


class TestNormalizeDataframe:
    def test_collision_does_not_raise(self):
        source = pd.DataFrame({"Name": ["a"], "Amount": ["$10"], "Total": ["$20"]})
        result = normalize_dataframe(source, "collide.csv")
        assert result["revenue"].iloc[0] == 10.0
        assert "revenue_2" in result.columns

    def test_strips_whitespace_and_tags_source(self):
        source = pd.DataFrame({"Name": ["  Sam Turner "], "Amount": ["$5,400.00"]})
        result = normalize_dataframe(source, "leads.csv")
        assert result["name"].iloc[0] == "Sam Turner"
        assert result["revenue"].iloc[0] == 5400.0
        assert result["source_file"].iloc[0] == "leads.csv"

    def test_drops_null_and_whitespace_only_rows(self):
        source = pd.DataFrame({
            "Name": ["a", None, "   ", " real "],
            "Amount": ["$1", None, "\t", "2"],
        })
        result = normalize_dataframe(source, "f.csv")
        assert result["name"].tolist() == ["a", "real"]

    def test_preserves_input_source_file_column(self):
        source = pd.DataFrame({"source_file": ["user value"], "Name": ["a"]})
        result = normalize_dataframe(source, "actual.csv")
        assert result["source_file_2"].iloc[0] == "user value"
        assert result["source_file"].iloc[0] == "actual.csv"

    @pytest.mark.parametrize("payload", ["=2+3", "+2+3", "-2+3", "@SUM(A1:A2)"])
    def test_neutralizes_spreadsheet_formula_prefixes(self, payload):
        result = normalize_dataframe(pd.DataFrame({"Name": [payload]}), "f.csv")
        assert result["name"].iloc[0] == "'" + payload

    def test_normalizes_dates_and_keeps_salary_separate(self):
        source = pd.DataFrame({"Hire Date": ["March 5, 2024"], "Salary": ["$72,000"]})
        result = normalize_dataframe(source, "hr.csv")
        assert str(result["hire_date"].iloc[0]) == "2024-03-05"
        assert result["salary"].iloc[0] == 72000.0
        assert "revenue" not in result.columns

    def test_keeps_invalid_money_text_visible_for_review(self):
        result = normalize_dataframe(
            pd.DataFrame({"Revenue": ["invoice 123"]}),
            "bad-money.csv",
        )
        assert result["revenue"].iloc[0] == "invoice 123"

    def test_does_not_mutate_input_dataframe(self):
        source = pd.DataFrame({"Name": ["  a  "]})
        normalize_dataframe(source, "f.csv")
        assert list(source.columns) == ["Name"]
        assert source.iloc[0, 0] == "  a  "


class TestLoadFile:
    def test_csv_preserves_leading_zero_identifier(self, tmp_path):
        path = tmp_path / "ids.csv"
        path.write_text("customer_id,name\n00123,Alice\n", encoding="utf-8")
        result = load_file(path)
        assert result is not None
        assert result["customer_id"].iloc[0] == "00123"

    def test_xlsx_formula_is_preserved_as_text_then_neutralized(self, tmp_path):
        path = tmp_path / "formula.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["Name"])
        sheet.append(["=1+1"])
        workbook.save(path)

        loaded = load_file(path)
        assert loaded is not None
        assert loaded["Name"].iloc[0] == "=1+1"
        normalized = normalize_dataframe(loaded, "formula.xlsx")
        assert normalized["name"].iloc[0] == "'=1+1"


class TestMergeFiles:
    @staticmethod
    def source_folder(tmp_path):
        source = tmp_path / "input"
        source.mkdir()
        return source

    def test_merges_deduplicates_and_writes_summary_first(self, tmp_path):
        source = self.source_folder(tmp_path)
        pd.DataFrame({"Name": ["a", "b"], "Sales Amount": ["$1", "$2"]}).to_csv(
            source / "one.csv", index=False
        )
        pd.DataFrame({"name": ["b", "c"], "amount": ["$2", "$3"]}).to_csv(
            source / "two.csv", index=False
        )
        output = tmp_path / "master.xlsx"

        summary = merge_files(source, output)

        workbook = load_workbook(output, read_only=True)
        assert workbook.sheetnames == ["Merge Summary", "Merged Data"]
        workbook.close()
        merged = pd.read_excel(output, sheet_name="Merged Data")
        assert len(merged) == 3
        assert list(merged.columns)[-1] == "source_file"
        assert summary["duplicates_removed"] == 1
        assert summary["final_rows"] == 3

    def test_no_dedup_keeps_duplicates(self, tmp_path):
        source = self.source_folder(tmp_path)
        for name in ("one.csv", "two.csv"):
            pd.DataFrame({"Name": ["a"], "Amount": ["$1"]}).to_csv(
                source / name, index=False
            )
        output = tmp_path / "master.xlsx"
        assert merge_files(source, output, deduplicate=False)["final_rows"] == 2

    def test_empty_folder_raises(self, tmp_path):
        source = self.source_folder(tmp_path)
        with pytest.raises(FileNotFoundError):
            merge_files(source, tmp_path / "master.xlsx")

    def test_finds_nested_case_insensitive_files_and_tracks_relative_source(self, tmp_path):
        source = self.source_folder(tmp_path)
        nested = source / "north"
        nested.mkdir()
        pd.DataFrame({"Name": ["a"]}).to_csv(nested / "DATA.CSV", index=False)
        output = tmp_path / "master.xlsx"

        merge_files(source, output)

        merged = pd.read_excel(output, sheet_name="Merged Data")
        assert merged["source_file"].tolist() == ["north/DATA.CSV"]

    def test_creates_missing_output_directories(self, tmp_path):
        source = self.source_folder(tmp_path)
        pd.DataFrame({"Name": ["a"]}).to_csv(source / "data.csv", index=False)
        output = tmp_path / "new" / "nested" / "master.xlsx"
        merge_files(source, output)
        assert output.exists()

    def test_rejects_output_inside_input_folder(self, tmp_path):
        source = self.source_folder(tmp_path)
        pd.DataFrame({"Name": ["a"]}).to_csv(source / "data.csv", index=False)
        with pytest.raises(ValueError, match="outside the input folder"):
            merge_files(source, source / "master.xlsx")

    def test_existing_output_requires_explicit_overwrite(self, tmp_path):
        source = self.source_folder(tmp_path)
        pd.DataFrame({"Name": ["a"]}).to_csv(source / "data.csv", index=False)
        output = tmp_path / "master.xlsx"
        merge_files(source, output)
        with pytest.raises(FileExistsError, match="--overwrite"):
            merge_files(source, output)
        assert merge_files(source, output, overwrite=True)["final_rows"] == 1

    def test_rejects_non_xlsx_output(self, tmp_path):
        source = self.source_folder(tmp_path)
        pd.DataFrame({"Name": ["a"]}).to_csv(source / "data.csv", index=False)
        with pytest.raises(ValueError, match=".xlsx"):
            merge_files(source, tmp_path / "master.csv")

    def test_partial_run_names_unreadable_files(self, tmp_path):
        source = self.source_folder(tmp_path)
        pd.DataFrame({"Name": ["a"]}).to_csv(source / "good.csv", index=False)
        (source / "bad.csv").write_bytes(b"\xff\xfe\x00\x00")
        output = tmp_path / "master.xlsx"

        summary = merge_files(source, output)

        assert summary["status"] == "partial"
        assert summary["files_skipped"] == ["bad.csv"]
        assert output.exists()

    def test_strict_mode_refuses_partial_output(self, tmp_path):
        source = self.source_folder(tmp_path)
        pd.DataFrame({"Name": ["a"]}).to_csv(source / "good.csv", index=False)
        (source / "bad.csv").write_bytes(b"\xff\xfe\x00\x00")
        output = tmp_path / "master.xlsx"

        with pytest.raises(ValueError, match="Strict mode"):
            merge_files(source, output, strict=True)
        assert not output.exists()

    def test_formula_payload_is_text_in_exported_workbook(self, tmp_path):
        source = self.source_folder(tmp_path)
        pd.DataFrame({"Name": ["=HYPERLINK(\"https://evil.test\")"]}).to_csv(
            source / "data.csv", index=False
        )
        output = tmp_path / "master.xlsx"
        merge_files(source, output)

        workbook = load_workbook(output, data_only=False)
        cell = workbook["Merged Data"]["A2"]
        assert cell.data_type != "f"
        assert cell.value.startswith("'=")
        workbook.close()

    def test_formula_like_filename_is_text_in_summary_sheet(self, tmp_path):
        source = self.source_folder(tmp_path)
        pd.DataFrame({"Name": ["safe"]}).to_csv(source / "=1+1.csv", index=False)
        output = tmp_path / "master.xlsx"

        merge_files(source, output)

        workbook = load_workbook(output, data_only=False)
        cell = workbook["Merge Summary"]["A10"]
        assert cell.data_type != "f"
        assert cell.value == "'=1+1.csv"
        workbook.close()

    def test_output_has_filter_table_freeze_and_number_formats(self, tmp_path):
        source = self.source_folder(tmp_path)
        pd.DataFrame({"Name": ["a"], "Revenue": ["$1.25"], "Date": ["2024/01/02"]}).to_csv(
            source / "data.csv", index=False
        )
        output = tmp_path / "master.xlsx"
        merge_files(source, output)

        workbook = load_workbook(output)
        sheet = workbook["Merged Data"]
        assert sheet.freeze_panes == "A2"
        assert len(sheet.tables) == 1
        headers = {cell.value: cell.column for cell in sheet[1]}
        assert "$" in sheet.cell(2, headers["revenue"]).number_format
        assert sheet.cell(2, headers["date"]).number_format == "yyyy-mm-dd"
        workbook.close()


def test_clean_string_passes_through_non_strings():
    assert clean_string(42) == 42
    assert clean_string("  x ") == "x"


def test_import_and_help_have_no_log_file_side_effect(tmp_path):
    environment = os.environ.copy()
    environment["PYTHONPATH"] = str(PROJECT_ROOT)
    subprocess.run(
        [sys.executable, "-W", "error", "-c", "import excel_merger"],
        cwd=tmp_path,
        env=environment,
        check=True,
        capture_output=True,
        text=True,
    )
    subprocess.run(
        [sys.executable, str(PROJECT_ROOT / "excel_merger.py"), "--help"],
        cwd=tmp_path,
        check=True,
        capture_output=True,
        text=True,
    )
    assert not (tmp_path / "merger.log").exists()


def test_sample_generator_creates_missing_destination_and_real_duplicate(tmp_path):
    destination = tmp_path / "new" / "sample_input"
    create_samples(destination)
    assert {path.name for path in destination.iterdir()} == {
        "hr_employees.csv", "marketing_leads.csv", "ops_data.xlsx",
        "sales_q1.xlsx", "sales_q2.xlsx",
    }
    summary = merge_files(destination, tmp_path / "sample_master.xlsx")
    assert summary["rows_loaded"] == 29
    assert summary["blank_rows_removed"] == 1
    assert summary["duplicates_removed"] == 1
    assert summary["final_rows"] == 27


def test_json_summary_creates_parent(tmp_path):
    path = tmp_path / "nested" / "summary.json"
    write_json_summary(path, {"final_rows": 3})
    assert json.loads(path.read_text(encoding="utf-8"))["final_rows"] == 3


def test_json_summary_cannot_replace_workbook(tmp_path, monkeypatch):
    source = tmp_path / "input"
    source.mkdir()
    pd.DataFrame({"Name": ["safe"]}).to_csv(source / "data.csv", index=False)
    output = tmp_path / "master.xlsx"
    monkeypatch.setattr(merger, "configure_logging", lambda *_args, **_kwargs: None)

    with pytest.raises(SystemExit) as exc_info:
        merger.main([
            "--input", str(source),
            "--output", str(output),
            "--json-out", str(output),
            "--overwrite",
        ])

    assert exc_info.value.code == 2
    assert not output.exists()


def test_json_summary_cannot_be_written_inside_input_folder(tmp_path):
    source = tmp_path / "input"
    source.mkdir()
    with pytest.raises(ValueError, match="outside the input folder"):
        validate_json_output_path(
            source,
            tmp_path / "master.xlsx",
            source / "run.json",
            overwrite=True,
        )


def test_existing_json_summary_requires_overwrite(tmp_path):
    source = tmp_path / "input"
    source.mkdir()
    json_path = tmp_path / "run.json"
    json_path.write_text("keep", encoding="utf-8")

    with pytest.raises(FileExistsError, match="--overwrite"):
        validate_json_output_path(
            source,
            tmp_path / "master.xlsx",
            json_path,
            overwrite=False,
        )
    assert json_path.read_text(encoding="utf-8") == "keep"


def test_main_writes_json_and_returns_partial_status(tmp_path, monkeypatch):
    source = tmp_path / "input"
    source.mkdir()
    pd.DataFrame({"Name": ["a"]}).to_csv(source / "good.csv", index=False)
    (source / "bad.csv").write_bytes(b"\xff\xfe\x00\x00")
    output = tmp_path / "master.xlsx"
    summary_path = tmp_path / "run.json"
    monkeypatch.setattr(merger, "configure_logging", lambda *_args, **_kwargs: None)

    result = merger.main([
        "--input", str(source),
        "--output", str(output),
        "--json-out", str(summary_path),
    ])

    assert result == 2
    assert json.loads(summary_path.read_text(encoding="utf-8"))["status"] == "partial"
