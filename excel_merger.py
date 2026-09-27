"""
Excel & CSV File Merger
=======================
Scans a folder for all Excel (.xlsx, .xls) and CSV files, intelligently
merges them into a single normalized master file, and outputs a clean Excel
file. Handles messy real-world data: different column orders, missing columns,
duplicate rows, inconsistent formatting, and varied column naming conventions.

Usage:
    python excel_merger.py --input ./sample_input --output merged_master.xlsx
    python excel_merger.py --input /path/to/folder --no-dedup

Features:
    - Discovers all .xlsx, .xls, and .csv files recursively
    - Normalizes column names (strips whitespace, lowercases for matching)
    - Maps common column aliases (e.g. "Sales Amount" → "revenue")
    - Fills missing columns with empty values
    - Deduplicates identical rows (configurable)
    - Strips leading/trailing whitespace from all text cells
    - Normalizes currency strings ("$5,400.00" → 5400.00)
    - Adds a "source_file" column tracking which file each row came from
    - Outputs a formatted Excel file with statistics summary sheet
"""

import argparse
import datetime as dt
import json
import logging
import math
import re
import tempfile
from logging.handlers import RotatingFileHandler
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

logger = logging.getLogger(__name__)

SUPPORTED_SUFFIXES = {".xlsx", ".xls", ".csv"}
SOURCE_COLUMN = "source_file"
MONEY_COLUMNS = {"revenue", "salary"}
DATE_COLUMNS = {"date", "hire_date"}


# ---------------------------------------------------------------------------
# Logging
# ---------------------------------------------------------------------------
def configure_logging(log_path: Path | None = Path("merger.log")) -> None:
    """Configure bounded runtime logging without import-time file I/O."""
    for handler in list(logger.handlers):
        logger.removeHandler(handler)
        handler.close()

    formatter = logging.Formatter("%(asctime)s [%(levelname)s] %(message)s")
    console = logging.StreamHandler()
    console.setFormatter(formatter)

    logger.setLevel(logging.INFO)
    logger.propagate = False
    logger.addHandler(console)

    if log_path is None:
        return
    try:
        log_path.parent.mkdir(parents=True, exist_ok=True)
        file_handler = RotatingFileHandler(
            log_path,
            maxBytes=1_000_000,
            backupCount=3,
            encoding="utf-8",
        )
        file_handler.setFormatter(formatter)
        logger.addHandler(file_handler)
    except OSError as exc:
        logger.warning("Could not write log file %s: %s", log_path, exc)

# ---------------------------------------------------------------------------
# Column alias map — maps known variants to a canonical name
# Add more entries here as you encounter new file formats.
# ---------------------------------------------------------------------------
COLUMN_ALIASES: dict[str, str] = {
    # Name variants
    "full name": "name",
    "full_name": "name",
    "first name": "first_name",
    "first_name": "first_name",
    "last name": "last_name",
    "last_name": "last_name",
    "employee name": "name",
    # Email variants
    "contact email": "email",
    "email address": "email",
    "e-mail": "email",
    # Revenue / money variants
    "revenue": "revenue",
    "sales amount": "revenue",
    "sales_amount": "revenue",
    "amount": "revenue",
    "salary": "salary",
    "total": "revenue",
    # Region / territory variants
    "region": "region",
    "territory": "region",
    "area": "region",
    "zone": "region",
    # Date variants
    "date": "date",
    "transaction date": "date",
    "transaction_date": "date",
    "hire date": "hire_date",
    "hire_date": "hire_date",
    "sale date": "date",
    # Department variants
    "department": "department",
    "dept": "department",
    "team": "department",
    "division": "department",
}


def normalize_column_name(col: str) -> str:
    """
    Lowercase, strip whitespace, and apply alias mapping.
    Returns the canonical column name.
    """
    cleaned = re.sub(r"\s+", " ", str(col).strip().casefold())
    canonical = COLUMN_ALIASES.get(cleaned, cleaned)
    return excel_safe_text(canonical)


def dedupe_columns(cols: list[str], source_filename: str) -> list[str]:
    """
    Ensure canonical column names are unique within a single file.

    Multiple source headers can alias to the same canonical name (e.g. both
    "Amount" and "Total" map to "revenue"). Left unhandled, pandas produces two
    identically-named columns, `df["revenue"]` returns a DataFrame instead of a
    Series, and every downstream scalar operation raises
    "The truth value of a Series is ambiguous".
    """
    used: set[str] = set()
    next_suffix: dict[str, int] = {}
    out: list[str] = []
    for c in cols:
        renamed = c
        if renamed in used:
            suffix = next_suffix.get(c, 2)
            renamed = f"{c}_{suffix}"
            while renamed in used:
                suffix += 1
                renamed = f"{c}_{suffix}"
            next_suffix[c] = suffix + 1
            logger.warning(
                f"{source_filename}: duplicate canonical column '{c}' renamed to "
                f"'{renamed}'. Review COLUMN_ALIASES if these should be merged."
            )
        used.add(renamed)
        next_suffix.setdefault(c, 2)
        out.append(renamed)
    return out


_CURRENCY_RE = re.compile(
    r"^(?P<sign>[+-]?)(?P<amount>(?:\d{1,3}(?:,\d{3})+|\d+)(?:\.\d+)?|\.\d+)$"
)


def normalize_currency(value) -> float | None:
    """
    Convert currency strings like '$5,400.00' or '5400' to float.
    Returns None if conversion is not possible.
    """
    if isinstance(value, (pd.Series, pd.DataFrame)):
        raise TypeError(
            "normalize_currency expects a scalar. Receiving a Series means the "
            "DataFrame has duplicate column names -- run dedupe_columns() first."
        )
    if pd.isna(value) or isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        numeric = float(value)
        return numeric if math.isfinite(numeric) else None

    raw = str(value).strip()
    negative_parentheses = raw.startswith("(") and raw.endswith(")")
    if negative_parentheses:
        raw = raw[1:-1].strip()
    raw = re.sub(r"(?i)\b(?:USD|EUR|GBP|CAD|AUD)\b", "", raw)
    raw = re.sub(r"[$€£¥\s]", "", raw)
    match = _CURRENCY_RE.fullmatch(raw)
    if not match:
        return None
    numeric = float(match.group("amount").replace(",", ""))
    if match.group("sign") == "-" or negative_parentheses:
        numeric = -numeric
    return numeric if math.isfinite(numeric) else None


def _normalize_money_cell(value: Any) -> Any:
    """Type valid money while keeping invalid source text visible for review."""
    normalized = normalize_currency(value)
    if normalized is not None or pd.isna(value):
        return normalized
    raw = clean_string(value)
    return None if raw == "" else excel_safe_text(raw)


def clean_string(value) -> str:
    """Strip whitespace from string values; leave non-strings unchanged."""
    if isinstance(value, str):
        return value.strip()
    return value


def excel_safe_text(value: Any) -> Any:
    """Prevent untrusted text from becoming an executable Excel formula."""
    if isinstance(value, str) and value.startswith(("=", "+", "-", "@", "\t", "\r")):
        return "'" + value
    return value


def normalize_date(value: Any) -> Any:
    """Return a typed date when parseable, preserving unrecognized text."""
    if pd.isna(value):
        return None
    if isinstance(value, (dt.datetime, dt.date, pd.Timestamp)):
        return pd.Timestamp(value).date()
    raw = str(value).strip()
    if not raw:
        return None
    parsed = pd.to_datetime(raw, errors="coerce", format="mixed")
    return raw if pd.isna(parsed) else parsed.date()


# ---------------------------------------------------------------------------
# File loading
# ---------------------------------------------------------------------------
def _read_xlsx(path: Path) -> pd.DataFrame:
    """Read the first worksheet while preserving formulas as inert text."""
    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        worksheet = workbook.active
        rows = worksheet.iter_rows(values_only=True)
        try:
            raw_headers = next(rows)
        except StopIteration:
            return pd.DataFrame()
        headers = [
            f"unnamed_{index + 1}"
            if value is None or not str(value).strip()
            else value
            for index, value in enumerate(raw_headers)
        ]
        return pd.DataFrame(rows, columns=headers)
    finally:
        workbook.close()


def load_file(path: Path) -> pd.DataFrame | None:
    """
    Load an Excel or CSV file into a DataFrame.
    Returns None if the file cannot be read.
    """
    try:
        suffix = path.suffix.lower()
        if suffix == ".xlsx":
            df = _read_xlsx(path)
        elif suffix == ".xls":
            df = pd.read_excel(path, engine="xlrd", dtype=object)
        elif suffix == ".csv":
            # Keep non-schema fields as text so identifiers such as 00123 are
            # not irreversibly converted to 123. Known money/date fields are
            # typed explicitly during normalization.
            df = pd.read_csv(path, dtype=object, encoding="utf-8-sig")
        else:
            logger.warning(f"Unsupported file type: {path.name}")
            return None

        logger.info(f"Loaded {path.name}: {len(df)} rows, {len(df.columns)} columns")
        return df

    except Exception as exc:
        logger.error("Failed to load %s: %s", path.name, exc)
        return None


# ---------------------------------------------------------------------------
# Per-file normalization
# ---------------------------------------------------------------------------
def normalize_dataframe(df: pd.DataFrame, source_filename: str) -> pd.DataFrame:
    """
    Apply all normalization steps to a single DataFrame:
    1. Normalize column names (strip + alias mapping)
    2. Strip whitespace from string cells
    3. Drop entirely empty rows
    4. Normalize configured money fields
    5. Normalize configured date fields
    6. Neutralize spreadsheet formula prefixes in untrusted text
    7. Add the source_file lineage column
    """
    # Work on a copy so a failed normalization never partially mutates the
    # caller's DataFrame.
    df = df.copy()

    # --- 1. Normalize column names ---
    # Reserve SOURCE_COLUMN so a user-provided column with that name is kept
    # under a collision-safe suffix instead of being silently overwritten.
    normalized_columns = [normalize_column_name(c) for c in df.columns]
    df.columns = dedupe_columns(
        [SOURCE_COLUMN, *normalized_columns], source_filename
    )[1:]

    # --- 2. Strip whitespace from all string cells ---
    df = df.map(clean_string)

    # --- 3. Drop rows empty after trimming ---
    empty_cells = df.isna() | df.eq("")
    df = df.loc[~empty_cells.all(axis=1)].copy()

    # --- 4. Normalize configured money columns ---
    for column in MONEY_COLUMNS & set(df.columns):
        df[column] = df[column].apply(_normalize_money_cell)

    # --- 5. Normalize known dates to typed cells ---
    for column in DATE_COLUMNS & set(df.columns):
        df[column] = df[column].apply(normalize_date)

    # --- 6. Neutralize formulas in every untrusted text cell ---
    df = df.map(excel_safe_text)

    # --- 7. Tag with source ---
    df[SOURCE_COLUMN] = excel_safe_text(source_filename)

    return df


# ---------------------------------------------------------------------------
# Merge engine
# ---------------------------------------------------------------------------
def merge_files(
    input_folder: Path,
    output_path: Path,
    deduplicate: bool = True,
    *,
    overwrite: bool = False,
    strict: bool = False,
) -> dict[str, Any]:
    """
    Main merge function. Discovers files, normalizes each, aligns columns,
    concatenates, deduplicates, and writes to Excel.
    """
    input_folder = input_folder.resolve()
    output_path = output_path.resolve()
    if not input_folder.is_dir():
        raise ValueError(f"Input path is not a directory: {input_folder}")
    if output_path.suffix.casefold() != ".xlsx":
        raise ValueError("Output path must use the .xlsx extension")
    try:
        output_path.relative_to(input_folder)
    except ValueError:
        pass
    else:
        raise ValueError(
            "Output must be outside the input folder so source files can never "
            "be overwritten or re-ingested"
        )
    if output_path.exists() and not overwrite:
        raise FileExistsError(
            f"Output already exists: {output_path}. Use --overwrite to replace it."
        )

    # --- Discover all supported files recursively ---
    found_files = sorted(
        (
            path
            for path in input_folder.rglob("*")
            if path.is_file()
            and path.suffix.casefold() in SUPPORTED_SUFFIXES
            and not path.name.startswith("~$")
            and path.resolve() != output_path
        ),
        key=lambda path: path.relative_to(input_folder).as_posix().casefold(),
    )

    if not found_files:
        raise FileNotFoundError(f"No Excel or CSV files found in: {input_folder}")

    discovered_names = [path.relative_to(input_folder).as_posix() for path in found_files]
    logger.info("Found %s file(s) to merge: %s", len(found_files), discovered_names)

    # --- Load and normalize each file ---
    frames: list[pd.DataFrame] = []
    load_stats: list[dict[str, Any]] = []
    failed_files: list[str] = []

    for path in found_files:
        source_filename = path.relative_to(input_folder).as_posix()
        df = load_file(path)
        if df is None:
            failed_files.append(source_filename)
            continue
        original_rows = len(df)
        df = normalize_dataframe(df, source_filename)
        frames.append(df)
        load_stats.append({
            "file": source_filename,
            "rows_loaded": original_rows,
            "rows_kept": len(df),
            "columns": len(df.columns) - 1,  # exclude source_file
        })

    if failed_files and strict:
        raise ValueError(
            "Strict mode refused a partial merge; unreadable files: "
            + ", ".join(failed_files)
        )
    if not frames:
        raise ValueError("No files could be successfully loaded")

    # --- Concatenate (pandas aligns on column names, fills missing with NaN) ---
    merged = pd.concat(frames, ignore_index=True, sort=False)
    rows_before_dedup = len(merged)
    logger.info(f"Combined total: {rows_before_dedup} rows")

    # --- Deduplicate (exclude source_file from key) ---
    dupes_removed = 0
    if deduplicate:
        key_cols = [c for c in merged.columns if c != SOURCE_COLUMN]
        merged.drop_duplicates(subset=key_cols, keep="first", inplace=True)
        merged.reset_index(drop=True, inplace=True)
        dupes_removed = rows_before_dedup - len(merged)
        if dupes_removed:
            logger.info(f"Removed {dupes_removed} duplicate row(s)")

    # --- Reorder: put source_file last ---
    cols = [c for c in merged.columns if c != SOURCE_COLUMN] + [SOURCE_COLUMN]
    merged = merged[cols]

    summary: dict[str, Any] = {
        "generated_at": dt.datetime.now(dt.UTC).isoformat(timespec="seconds"),
        "status": "partial" if failed_files else "complete",
        "input_folder": str(input_folder),
        "output_path": str(output_path),
        "files_discovered": len(found_files),
        "files_merged": len(frames),
        "files_skipped": failed_files,
        "rows_loaded": sum(item["rows_loaded"] for item in load_stats),
        "rows_after_cleaning": rows_before_dedup,
        "blank_rows_removed": sum(
            item["rows_loaded"] - item["rows_kept"] for item in load_stats
        ),
        "duplicates_removed": dupes_removed,
        "final_rows": len(merged),
        "columns": list(merged.columns),
        "source_files": load_stats,
    }

    # --- Write output atomically ---
    write_excel(merged, output_path, load_stats, summary)
    logger.info(
        "Merge complete: %s file(s), %s rows written to %s",
        len(frames),
        len(merged),
        output_path,
    )
    return summary


# ---------------------------------------------------------------------------
# Excel writer with formatting
# ---------------------------------------------------------------------------
def write_excel(
    df: pd.DataFrame,
    output_path: Path,
    stats: list[dict[str, Any]],
    summary: dict[str, Any],
) -> None:
    """Atomically write a formatted summary and merged-data workbook."""
    output_path.parent.mkdir(parents=True, exist_ok=True)
    summary_df = pd.DataFrame([
        {
            # Filenames are untrusted input too. Keep a formula-like filename
            # visible while ensuring Excel stores it as inert text.
            "Source File": excel_safe_text(item["file"]),
            "Rows Loaded": item["rows_loaded"],
            "Rows Kept": item["rows_kept"],
            "Blank Rows Removed": item["rows_loaded"] - item["rows_kept"],
            "Original Columns": item["columns"],
        }
        for item in stats
    ])

    temp_path: Path | None = None
    try:
        with tempfile.NamedTemporaryFile(
            dir=output_path.parent,
            prefix=f".{output_path.stem}-",
            suffix=".xlsx",
            delete=False,
        ) as temp_file:
            temp_path = Path(temp_file.name)

        with pd.ExcelWriter(temp_path, engine="openpyxl") as writer:
            # Put the reader-facing summary first, followed by the audit data.
            summary_df.to_excel(
                writer,
                index=False,
                sheet_name="Merge Summary",
                startrow=8,
            )
            df.to_excel(writer, index=False, sheet_name="Merged Data")
            _format_summary_sheet(writer.sheets["Merge Summary"], stats, summary)
            _format_data_sheet(writer.sheets["Merged Data"], df)

        temp_path.replace(output_path)
    finally:
        if temp_path is not None:
            temp_path.unlink(missing_ok=True)


def _add_table(ws, ref: str, name: str) -> None:
    table = Table(displayName=name, ref=ref)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    ws.add_table(table)


def _format_summary_sheet(
    ws,
    stats: list[dict[str, Any]],
    summary: dict[str, Any],
) -> None:
    """Create a compact reader-facing summary above the source audit table."""
    dark = "1A3A5C"
    accent = "DCEAF7"
    muted = "5B7083"
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = dark

    ws["A2"] = "File merge summary"
    ws["A2"].font = Font(name="Arial", size=16, bold=True, color="172B3A")
    ws["A3"] = (
        f"{summary['files_merged']} source files normalized into one workbook. "
        f"Status: {summary['status']}."
    )
    ws["A3"].font = Font(name="Arial", size=10, italic=True, color=muted)

    metrics = [
        ("Files merged", summary["files_merged"]),
        ("Rows loaded", summary["rows_loaded"]),
        ("Blank rows removed", summary["blank_rows_removed"]),
        ("Duplicates removed", summary["duplicates_removed"]),
        ("Final rows", summary["final_rows"]),
        ("Files skipped", len(summary["files_skipped"])),
    ]
    for column, (label, value) in zip((1, 3, 5, 7, 9, 11), metrics, strict=True):
        label_cell = ws.cell(row=5, column=column, value=label)
        value_cell = ws.cell(row=6, column=column, value=value)
        label_cell.font = Font(name="Arial", size=9, bold=True, color=muted)
        value_cell.font = Font(name="Arial", size=16, bold=True, color=dark)
        value_cell.fill = PatternFill("solid", fgColor=accent)
        label_cell.alignment = value_cell.alignment = Alignment(vertical="center")

    ws["A8"] = "Source files"
    ws["A8"].font = Font(name="Arial", size=11, bold=True, color="172B3A")
    header_row = 9
    last_row = header_row + len(stats)
    last_column = 5
    if stats:
        _add_table(
            ws,
            f"A{header_row}:{get_column_letter(last_column)}{last_row}",
            "SourceFilesTable",
        )

    for row in ws.iter_rows(min_row=header_row, max_row=max(last_row, header_row)):
        for cell in row:
            cell.font = Font(
                name="Arial",
                size=10,
                bold=cell.row == header_row,
                color="FFFFFF" if cell.row == header_row else "172B3A",
            )
            cell.alignment = Alignment(
                horizontal="center" if cell.row == header_row else "left",
                vertical="center",
            )
    for cell in ws[header_row]:
        if cell.column <= last_column:
            cell.fill = PatternFill("solid", fgColor=dark)

    widths = {"A": 38, "B": 15, "C": 14, "D": 20, "E": 18}
    for letter, width in widths.items():
        ws.column_dimensions[letter].width = width
    for letter in ("G", "I", "K"):
        ws.column_dimensions[letter].width = 18
    for letter in ("F", "H", "J"):
        ws.column_dimensions[letter].width = 3
    ws.row_dimensions[2].height = 24
    ws.row_dimensions[6].height = 26


def _format_data_sheet(ws, df: pd.DataFrame) -> None:
    """Format the normalized audit data for filtering and review."""
    dark = "1A3A5C"
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = "4F81BD"
    ws.freeze_panes = "A2"

    max_row = max(ws.max_row, 1)
    max_column = max(ws.max_column, 1)
    ref = f"A1:{get_column_letter(max_column)}{max_row}"
    if len(df) > 0:
        _add_table(ws, ref, "MergedDataTable")
    else:
        ws.auto_filter.ref = ref

    for cell in ws[1]:
        cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=dark)
        cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 24

    for row in ws.iter_rows(min_row=2):
        ws.row_dimensions[row[0].row].height = 20
        for cell in row:
            cell.font = Font(name="Arial", size=10, color="172B3A")
            cell.alignment = Alignment(vertical="center")

    headers = {cell.value: cell.column for cell in ws[1]}
    for column in MONEY_COLUMNS & set(headers):
        for cell in ws.iter_cols(
            min_col=headers[column],
            max_col=headers[column],
            min_row=2,
            max_row=max_row,
        ):
            for item in cell:
                item.number_format = '"$"#,##0.00;[Red]-"$"#,##0.00'
    for column in DATE_COLUMNS & set(headers):
        for cell in ws.iter_cols(
            min_col=headers[column],
            max_col=headers[column],
            min_row=2,
            max_row=max_row,
        ):
            for item in cell:
                item.number_format = "yyyy-mm-dd"

    for index, column_name in enumerate(df.columns, start=1):
        values = [column_name, *df[column_name].head(200).tolist()]
        max_length = max(len(str(value)) for value in values if value is not None)
        width = min(max(max_length + 2, 12), 38)
        if column_name == SOURCE_COLUMN:
            width = max(width, 24)
        ws.column_dimensions[get_column_letter(index)].width = width


def write_json_summary(path: Path, summary: dict[str, Any]) -> None:
    """Atomically write a machine-readable audit artifact."""
    path.parent.mkdir(parents=True, exist_ok=True)
    temp_path: Path | None = None
    try:
        with tempfile.NamedTemporaryFile(
            dir=path.parent,
            prefix=f".{path.stem}-",
            suffix=".json",
            delete=False,
        ) as temp_file:
            temp_path = Path(temp_file.name)
        with temp_path.open("w", encoding="utf-8") as handle:
            json.dump(summary, handle, indent=2, ensure_ascii=False)
        temp_path.replace(path)
    finally:
        if temp_path is not None:
            temp_path.unlink(missing_ok=True)


def validate_json_output_path(
    input_folder: Path,
    output_path: Path,
    json_path: Path,
    *,
    overwrite: bool,
) -> None:
    """Ensure an optional JSON sidecar cannot overwrite sources or the workbook."""
    input_folder = input_folder.resolve()
    output_path = output_path.resolve()
    json_path = json_path.resolve()

    if json_path == output_path:
        raise ValueError("JSON summary path must differ from the output workbook")
    if json_path.suffix.casefold() != ".json":
        raise ValueError("JSON summary path must use the .json extension")
    try:
        json_path.relative_to(input_folder)
    except ValueError:
        pass
    else:
        raise ValueError(
            "JSON summary must be outside the input folder so source files "
            "cannot be overwritten"
        )
    if json_path.exists() and not overwrite:
        raise FileExistsError(
            f"JSON summary already exists: {json_path}. "
            "Use --overwrite to replace it."
        )


# ---------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------
def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Excel & CSV File Merger")
    parser.add_argument(
        "--input",
        default="./sample_input",
        help="Folder containing Excel/CSV files to merge (default: ./sample_input)",
    )
    parser.add_argument(
        "--output",
        default="merged_master.xlsx",
        help="Output Excel filename (default: merged_master.xlsx)",
    )
    parser.add_argument(
        "--no-dedup",
        action="store_true",
        help="Skip duplicate row removal",
    )
    parser.add_argument(
        "--overwrite",
        action="store_true",
        help="Replace existing output artifacts (source files are never overwritten)",
    )
    parser.add_argument(
        "--strict",
        action="store_true",
        help="Abort instead of creating a partial workbook if any input cannot be read",
    )
    parser.add_argument(
        "--json-out",
        help="Optional path for a machine-readable run summary",
    )
    args = parser.parse_args(argv)

    input_folder = Path(args.input).resolve()
    output_path = Path(args.output).resolve()
    json_output_path = Path(args.json_out).resolve() if args.json_out else None
    configure_logging(Path("merger.log").resolve())

    try:
        if json_output_path is not None:
            validate_json_output_path(
                input_folder,
                output_path,
                json_output_path,
                overwrite=args.overwrite,
            )
        summary = merge_files(
            input_folder=input_folder,
            output_path=output_path,
            deduplicate=not args.no_dedup,
            overwrite=args.overwrite,
            strict=args.strict,
        )
        if json_output_path is not None:
            write_json_summary(json_output_path, summary)
    except (FileExistsError, FileNotFoundError, OSError, ValueError) as exc:
        parser.error(str(exc))

    label = "OK" if summary["status"] == "complete" else "PARTIAL"
    print(
        f"\n[{label}] Merged {summary['files_merged']} file(s) -> "
        f"{summary['final_rows']} rows -> {output_path}"
    )
    if summary["files_skipped"]:
        print("Skipped unreadable files: " + ", ".join(summary["files_skipped"]))
    return 0 if summary["status"] == "complete" else 2


if __name__ == "__main__":
    raise SystemExit(main())
